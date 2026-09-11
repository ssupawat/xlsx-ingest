"""Files, sheets and the real insert path, ported from stress_test.py."""
import zipfile

import pytest
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.exc import NoSuchTableError

from excel_to_sql import IngestionError, excel_to_sql

from sheets import HEADER, OK


def test_totally_empty_sheet_rejected(make_sheet, engine):
    # Arrange
    p = make_sheet("empty", [], header=None)

    # Act / Assert
    with pytest.raises(ValueError, match="is empty"):
        excel_to_sql(p, "orders", engine, dry_run=True)


def test_first_sheet_is_the_default(make_sheet, engine):
    # Arrange
    p = make_sheet("multi", [OK], sheets=["Data", "Summary"])

    # Act
    n = excel_to_sql(p, "orders", engine, dry_run=True)

    # Assert
    assert n == 1


def test_sheet_by_name(make_sheet, engine):
    # Arrange: Summary carries a header only
    p = make_sheet("multi", [OK], sheets=["Data", "Summary"])

    # Act
    n = excel_to_sql(p, "orders", engine, dry_run=True, sheet_name="Summary")

    # Assert
    assert n == 0


def test_missing_sheet_name_raises(make_sheet, engine):
    # Arrange
    p = make_sheet("multi", [OK], sheets=["Data", "Summary"])

    # Act / Assert
    with pytest.raises(Exception, match="Nope"):
        excel_to_sql(p, "orders", engine, dry_run=True, sheet_name="Nope")


def test_lower_cells_of_a_merged_range_rejected(tmp_path, engine):
    # Arrange
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append(OK)
    ws.append([2, "ACME", 3, None, None, None, None, None, None])
    ws.merge_cells("B2:B3")
    p = tmp_path / "merged.xlsx"
    wb.save(p)

    # Act
    with pytest.raises(IngestionError) as e:
        excel_to_sql(p, "orders", engine, dry_run=True)

    # Assert
    assert "[Sheet!B3]" in str(e.value)
    assert "part of merged range at B2" in str(e.value)


def test_formula_without_cached_result_rejected(tmp_path, engine):
    # Arrange: openpyxl writes no <v> for the formula
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append([1, "ACME", "=5+5", None, None, None, None, None, None])
    p = tmp_path / "formula.xlsx"
    wb.save(p)

    # Act
    with pytest.raises(IngestionError) as e:
        excel_to_sql(p, "orders", engine, dry_run=True)

    # Assert
    assert "[Sheet!C2]" in str(e.value)
    assert "formula with no cached result" in str(e.value)


def test_cached_formula_passes_but_error_cell_rejected(tmp_path, engine):
    # Arrange: patch a cached formula into C2 and an #N/A error cell into C3
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append([1, "ACME", 1, None, None, None, None, None, "z"])
    ws.append([2, "ACME", 1, None, None, None, None, None, "z"])
    src = tmp_path / "f_base.xlsx"
    wb.save(src)
    p = tmp_path / "f_patched.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(p, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                x = data.decode()
                x = x.replace('<c r="C2" t="n"><v>1</v></c>',
                              '<c r="C2"><f>2+3</f><v>5</v></c>')
                x = x.replace('<c r="C3" t="n"><v>1</v></c>',
                              '<c r="C3" t="e"><f>NA()</f><v>#N/A</v></c>')
                data = x.encode()
            zout.writestr(item, data)

    # Act
    with pytest.raises(IngestionError) as e:
        excel_to_sql(p, "orders", engine, dry_run=True)

    # Assert
    assert "1 cell(s) rejected" in str(e.value)
    assert "[Sheet!C3]" in str(e.value)
    assert "Excel error value #N/A" in str(e.value)


def test_duplicate_pk_insert_rolls_back(engine, make_sheet):
    # Arrange
    p = make_sheet("duppk", [OK, OK])

    # Act
    with pytest.raises(Exception, match="UNIQUE constraint failed"):
        excel_to_sql(p, "orders", engine)

    # Assert: the single transaction left nothing behind
    with engine.connect() as conn:
        assert conn.execute(text("select count(*) from orders")).scalar() == 0


def test_unknown_table_raises(make_sheet, engine):
    # Arrange
    p = make_sheet("any", [OK])

    # Act / Assert
    with pytest.raises(NoSuchTableError):
        excel_to_sql(p, "nope", engine, dry_run=True)


def test_csv_rejected(tmp_path, engine):
    # Arrange
    p = tmp_path / "data.csv"
    p.write_text("id,customer,qty\n1,ACME,10\n")

    # Act / Assert
    with pytest.raises(ValueError, match="only .xlsx is accepted"):
        excel_to_sql(p, "orders", engine, dry_run=True)


def test_lying_display_formats_rejected(tmp_path, engine):
    # Arrange: five formats whose display differs from the stored value,
    # two rows whose formats are honest (separators only, General)
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for i, fmt in enumerate(["0%", '"฿"#,##0.00', "0.00E+00", ";;;", "# ?/?"], start=1):
        ws.append([i, "A", 1, 0.15, None, None, None, None, None])
        ws.cell(row=i + 1, column=4).number_format = fmt
    ws.append([6, "A", 1, 1234.5, None, None, None, None, None])
    ws.cell(row=7, column=4).number_format = "#,##0.00"
    ws.append([7, "A", 1, 1234.5, None, None, None, None, None])
    p = tmp_path / "formats.xlsx"
    wb.save(p)

    # Act
    with pytest.raises(IngestionError) as e:
        excel_to_sql(p, "orders", engine, dry_run=True)

    # Assert
    msg = str(e.value)
    assert "5 cell(s) rejected in 5 distinct problem(s)" in msg
    for coord in ("D2", "D3", "D4", "D5", "D6"):
        assert f"[Sheet!{coord}]" in msg
    assert "D7" not in msg and "D8" not in msg
    assert "percent format" in msg
    assert "currency format" in msg
    assert "scientific format" in msg
    assert "fraction format" in msg


def test_percent_format_on_text_column_rejected(tmp_path, engine):
    # Arrange
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append([1, "A", 1, None, None, None, None, None, 0.5])
    ws.cell(row=2, column=9).number_format = "0%"
    p = tmp_path / "pcttext.xlsx"
    wb.save(p)

    # Act
    with pytest.raises(IngestionError) as e:
        excel_to_sql(p, "orders", engine, dry_run=True)

    # Assert
    assert "[Sheet!I2]" in str(e.value)
    assert "percent format" in str(e.value)
