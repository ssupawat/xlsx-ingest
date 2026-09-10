import datetime as dt, zipfile
from pathlib import Path
from openpyxl import Workbook
H=["id","customer","qty","price","ship_date","created_at","pickup","paid","note"]
wb=Workbook(); ws=wb.active; ws.title="Data"; ws.append(H)
ws.append([1,"ACME",1,0.15,dt.date(2026,1,5),None,None,True,"a"])   # percent
ws.cell(row=2,column=4).number_format='0%'
ws.append([2,"ACME",1,1234.5,dt.date(2026,1,6),None,None,True,"b"]) # currency
ws.cell(row=3,column=4).number_format='"฿"#,##0.00'
ws.append([3,"MERGE",1,1,dt.date(2026,1,7),None,None,True,"c"])
ws.append([4,None,1,1,dt.date(2026,1,8),None,None,True,"d"])
ws.merge_cells("B4:B5")
ws.append([5,"ACME",1,1,dt.date(2026,1,9),None,None,True,"e"])      # -> #N/A qty
ws.append([6,"ACME",1,1,dt.date(2026,1,10),None,None,True,"f"])     # -> formula no cache
wb.save("rt_base.xlsx")
src=Path("rt_base.xlsx"); dst=Path("rt_orig.xlsx")
with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst,"w") as zo:
    for it in zi.infolist():
        d=zi.read(it.filename)
        if it.filename=="xl/worksheets/sheet1.xml":
            x=d.decode()
            x=x.replace('<c r="C6" t="n"><v>1</v></c>','<c r="C6" t="e"><f>NA()</f><v>#N/A</v></c>')
            x=x.replace('<c r="C7" t="n"><v>1</v></c>','<c r="C7"><f>1+0</f><v></v></c>')
            d=x.encode()
        zo.writestr(it,d)
print("built rt_orig.xlsx")
